from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
from docx import Document


RAW_DIR = Path("data/raw/theory")

OUTPUT_DIR = Path("data/processed/extracted")


def extract_pdf(file_path):
    reader = PdfReader(file_path)

    text = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text.append(page_text)

    return "\n".join(text)


def extract_pptx(file_path):
    presentation = Presentation(file_path)

    text = []

    for slide in presentation.slides:
        for shape in slide.shapes:

            if hasattr(shape, "text"):
                if shape.text.strip():
                    text.append(shape.text)

    return "\n".join(text)


def extract_docx(file_path):
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text)


def extract_text(file_path):

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    elif extension in [".ppt", ".pptx"]:
        return extract_pptx(file_path)

    elif extension == ".docx":
        return extract_docx(file_path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")


if __name__ == "__main__":

    print("OSTutorLLM - Document Ingestion")
    print("--------------------------------")

    print(f"Looking for documents in: {RAW_DIR}")

    if not RAW_DIR.exists():
        print("Theory folder does not exist yet.")
        print("Please extract Theory.zip into data/raw/theory/")
        exit()

    # Create the output directory if it doesn't exist
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    files = [
        file
        for file in RAW_DIR.rglob("*")
        if file.suffix.lower() in [".pdf", ".pptx", ".docx"]
    ]

    print(f"Found {len(files)} documents.\n")

    successful = 0

    for file in files:

        print(f"Reading: {file.name}")

        try:
            text = extract_text(file)

            # Add the original file type to avoid name collisions
            output_name = f"{file.stem}__{file.suffix.lower().replace('.', '')}.txt"
            output_path = OUTPUT_DIR / output_name

            # Save extracted text
            output_path.write_text(text, encoding="utf-8")

            print(f"Extracted {len(text)} characters")
            print(f"Saved to: {output_path}\n")

            successful += 1

        except Exception as e:
            print(f"ERROR: Could not process {file.name}")
            print(f"Reason: {e}\n")

    print("--------------------------------")
    print(f"Successfully processed: {successful}/{len(files)} documents")
    print(f"Extracted text saved in: {OUTPUT_DIR}")